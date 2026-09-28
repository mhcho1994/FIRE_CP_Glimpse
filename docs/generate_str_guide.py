#!/usr/bin/env python3
"""Build an STR integration guide from a local Markdown source.

Adapted from the supplied Moonshot Classifier generator. Tool-specific title,
footer, metadata and default paths are no longer hard-coded. A Word template is
optional. Supported Markdown: headings, paragraphs, fenced code, pipe tables,
notes, lists, local images, links, inline code/bold and explicit page breaks.
Missing images and malformed fenced blocks fail the build by default.
"""
from __future__ import annotations

import argparse
import html
import sys
from dataclasses import dataclass
from datetime import datetime, timezone
import re
import tempfile
import zipfile
from pathlib import Path
from xml.etree import ElementTree

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.shared import Cm, Inches, Pt, RGBColor


@dataclass(frozen=True)
class GuideMetadata:
    title: str
    footer: str
    subject: str = "STR Integration Documentation"
    author: str = "FIRE / STR"
    keywords: str = "FIRE, STR, integration, documentation"


BLUE = "1F4E79"
MID_BLUE = "4F81BD"
LIGHT_BLUE = "D9EAF7"
CODE_BLUE = "EAF2F8"
LIGHT_GRAY = "F2F2F2"
WHITE = "FFFFFF"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=90, start=110, bottom=90, end=110) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def prevent_row_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def clear_paragraph(paragraph) -> None:
    p = paragraph._element
    for child in list(p):
        p.remove(child)


def add_page_number(run) -> None:
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend((begin, instr, end))


def sanitize_template(document: Document, metadata: GuideMetadata) -> None:
    """Retain layout/styles but remove source body, images, headers and metadata."""
    document.styles["Footer"].paragraph_format.tab_stops.clear_all()
    body = document._element.body
    for child in list(body):
        if child.tag != qn("w:sectPr"):
            body.remove(child)
    for rel_id, rel in list(document.part.rels.items()):
        if rel.reltype in (RT.IMAGE, RT.HYPERLINK):
            del document.part.rels[rel_id]
    for section in document.sections:
        # Unlink before clearing so an earlier section cannot retain stale text.
        for container in (section.header, section.first_page_header,
                          section.even_page_header, section.footer,
                          section.first_page_footer, section.even_page_footer):
            container.is_linked_to_previous = False
            for child in list(container._element):
                container._element.remove(child)
            container.add_paragraph()
            for rel_id, rel in list(container.part.rels.items()):
                if rel.reltype in (RT.IMAGE, RT.HYPERLINK):
                    del container.part.rels[rel_id]
        section.different_first_page_header_footer = False
        footer = section.footer.paragraphs[0]
        footer.style = document.styles["Footer"]
        footer.alignment = WD_ALIGN_PARAGRAPH.LEFT
        usable_width = section.page_width - section.left_margin - section.right_margin
        footer.paragraph_format.tab_stops.add_tab_stop(usable_width, WD_TAB_ALIGNMENT.RIGHT)
        for text in (metadata.footer, "\t"):
            run = footer.add_run(text)
            run.font.size = Pt(8)
            run.font.color.rgb = RGBColor.from_string(BLUE)
        page_run = footer.add_run()
        page_run.font.size = Pt(8)
        add_page_number(page_run)
    document.settings.odd_and_even_pages_header_footer = False
    props = document.core_properties
    props.title = metadata.title
    props.subject = metadata.subject
    props.author = metadata.author
    props.last_modified_by = metadata.author
    props.keywords = metadata.keywords
    props.category = "STR Integration Documentation"
    props.comments = "Generated from Markdown; edit the Markdown source, not this file."
    props.created = props.modified = datetime.now(timezone.utc)
    props.revision = 1
    props.identifier = ""
    props.language = "en-US"


def configure_styles(document: Document) -> None:
    normal = document.styles["Normal"]
    normal.font.name = "Liberation Sans"
    normal.font.size = Pt(10)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.10

    title = document.styles["Title"]
    title.font.name = "Liberation Sans"
    title.font.size = Pt(18)
    title.font.bold = True
    title.font.color.rgb = RGBColor.from_string(BLUE)
    title.paragraph_format.space_after = Pt(8)

    for style_name, size, color in (
        ("Heading 1", 14, BLUE),
        ("Heading 2", 12, "444444"),
        ("Heading 3", 10.5, MID_BLUE),
        ("Heading 4", 10, BLUE),
        ("Heading 5", 10, BLUE),
        ("Heading 6", 10, BLUE),
    ):
        style = document.styles[style_name]
        style.font.name = "Liberation Sans"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.keep_with_next = True
        if style_name == "Heading 1":
            style.paragraph_format.space_before = Pt(14)

    if "Code Block" not in [style.name for style in document.styles]:
        code = document.styles.add_style("Code Block", WD_STYLE_TYPE.PARAGRAPH)
    else:
        code = document.styles["Code Block"]
    code.font.name = "Liberation Mono"
    code.font.size = Pt(8.8)
    code.paragraph_format.space_after = Pt(0)
    code.paragraph_format.line_spacing = 1.0

    if "Document Subtitle" not in [style.name for style in document.styles]:
        subtitle = document.styles.add_style("Document Subtitle", WD_STYLE_TYPE.PARAGRAPH)
    else:
        subtitle = document.styles["Document Subtitle"]
    subtitle.font.name = "Liberation Sans"
    subtitle.font.size = Pt(10.5)
    subtitle.font.color.rgb = RGBColor.from_string(MID_BLUE)
    subtitle.paragraph_format.space_after = Pt(8)

    for name in ("List Bullet", "List Number"):
        document.styles[name].font.name = "Liberation Sans"
        document.styles[name].font.size = Pt(10)
        document.styles[name].paragraph_format.space_after = Pt(3)
    if "Ordered List" not in document.styles:
        document.styles.add_style("Ordered List", WD_STYLE_TYPE.PARAGRAPH)
    ordered = document.styles["Ordered List"]
    ordered.base_style = normal
    ordered.paragraph_format.left_indent = Inches(0.22)
    ordered.paragraph_format.first_line_indent = Inches(-0.22)
    ordered.paragraph_format.space_after = Pt(3)
    caption = document.styles["Caption"]
    caption.font.name = "Liberation Sans"
    caption.font.size = Pt(8.5)
    caption.font.color.rgb = RGBColor.from_string("555555")
    caption.paragraph_format.keep_with_next = False


def add_hyperlink(paragraph, text: str, url: str):
    part = paragraph.part
    rel_id = part.relate_to(url, RT.HYPERLINK, is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rel_id)
    run = OxmlElement("w:r")
    run_props = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), MID_BLUE)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    run_props.extend((color, underline))
    run.append(run_props)
    text_node = OxmlElement("w:t")
    text_node.text = text
    run.append(text_node)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


INLINE_PATTERN = re.compile(
    r"(\*\*.+?\*\*|`[^`]+`|\[[^\]]+\]\([^)]+\))"
)


def add_inline(paragraph, text: str, *, base_bold: bool = False) -> None:
    # The intentionally small syntax subset matches the supplied source guides.
    # A <br> in a table cell is a real line break, never literal HTML in Word.
    parts = re.split(r"(<br\s*/?>)", text, flags=re.IGNORECASE)
    for part in parts:
        if re.fullmatch(r"<br\s*/?>", part, flags=re.IGNORECASE):
            paragraph.add_run().add_break()
            continue
        position = 0
        for match in INLINE_PATTERN.finditer(part):
            if match.start() > position:
                run = paragraph.add_run(html.unescape(part[position:match.start()]))
                run.bold = base_bold
            token = match.group(0)
            if token.startswith("**"):
                run = paragraph.add_run(html.unescape(token[2:-2]))
                run.bold = True
            elif token.startswith("`"):
                run = paragraph.add_run(token[1:-1])
                run.font.name = "Liberation Mono"
                run.font.size = Pt(8.8)
                run.font.color.rgb = RGBColor.from_string(BLUE)
                run.bold = base_bold
            else:
                link = re.fullmatch(r"\[([^\]]+)\]\(([^)]+)\)", token)
                if link is None:
                    raise ValueError(f"Invalid inline link: {token}")
                add_hyperlink(paragraph, link.group(1), link.group(2))
            position = match.end()
        if position < len(part):
            run = paragraph.add_run(html.unescape(part[position:]))
            run.bold = base_bold


def split_markdown_row(line: str) -> list[str]:
    content = line.strip()
    if content.startswith("|"):
        content = content[1:]
    if content.endswith("|") and not content.endswith("\\|"):
        content = content[:-1]
    cells: list[str] = []
    current: list[str] = []
    index = 0
    while index < len(content):
        char = content[index]
        if char == chr(92) and index + 1 < len(content) and content[index + 1] in ("|", chr(92)):
            current.append(content[index + 1])
            index += 2
            continue
        if char == "|":
            cells.append("".join(current).strip())
            current = []
        else:
            current.append(char)
        index += 1
    cells.append("".join(current).strip())
    return cells


def is_table_separator(line: str) -> bool:
    cells = split_markdown_row(line)
    return bool(cells) and all(re.fullmatch(r":?-{3,}:?", cell) for cell in cells)


def add_markdown_table(document: Document, lines: list[str]) -> None:
    rows = [split_markdown_row(line) for line in lines]
    if len(rows) < 2 or not is_table_separator(lines[1]):
        raise ValueError("A pipe table needs a header separator row.")
    rows.pop(1)
    column_count = len(rows[0])
    if any(len(row) != column_count for row in rows):
        raise ValueError("Inconsistent column counts in Markdown table.")
    table = document.add_table(rows=len(rows), cols=column_count)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    section = document.sections[-1]
    width = section.page_width - section.left_margin - section.right_margin
    if column_count == 2:
        fractions = [0.28, 0.72]
    elif column_count == 3:
        fractions = [0.28, 0.36, 0.36]
    elif column_count == 4 and "Scenario connection" in rows[0]:
        fractions = [0.15, 0.13, 0.50, 0.22]
    elif column_count == 4:
        fractions = [0.13, 0.34, 0.24, 0.29]
    else:
        fractions = [1.0 / column_count] * column_count
    for col, fraction in zip(table.columns, fractions):
        col.width = int(width * fraction)
    for row_index, (word_row, values) in enumerate(zip(table.rows, rows)):
        prevent_row_split(word_row)
        if row_index == 0:
            repeat_table_header(word_row)
        for col_index, cell in enumerate(word_row.cells):
            cell.width = int(width * fractions[col_index])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
            set_cell_margins(cell, top=75, bottom=75)
            if row_index == 0:
                set_cell_shading(cell, BLUE)
            elif row_index % 2 == 0:
                set_cell_shading(cell, LIGHT_GRAY)
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1.05
            add_inline(paragraph, values[col_index], base_bold=row_index == 0)
            for run in paragraph.runs:
                run.font.size = Pt(8.6 if run.font.name == "Liberation Mono" else 9)
                if row_index == 0:
                    run.font.color.rgb = RGBColor.from_string(WHITE)
    after = document.add_paragraph()
    after.paragraph_format.space_after = Pt(0)
    after.paragraph_format.space_before = Pt(0)
    after.paragraph_format.line_spacing = 1
    after.add_run().font.size = Pt(2)


def add_code_block(document: Document, code: str) -> None:
    table = document.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    section = document.sections[-1]
    width = section.page_width - section.left_margin - section.right_margin
    table.columns[0].width = width
    cell = table.cell(0, 0)
    cell.width = width
    set_cell_shading(cell, CODE_BLUE)
    set_cell_margins(cell, top=95, start=130, bottom=95, end=130)
    lines = code.rstrip("\n").split("\n") or [""]
    if len(lines) <= 12:
        prevent_row_split(table.rows[0])
    paragraph = cell.paragraphs[0]
    paragraph.style = document.styles["Code Block"]
    paragraph.paragraph_format.widow_control = False
    for index, line in enumerate(lines):
        if index:
            paragraph.add_run().add_break()
        run = paragraph.add_run(line)
        run.font.name = "Liberation Mono"
        run.font.size = Pt(8.8)
    after = document.add_paragraph()
    after.paragraph_format.space_after = Pt(1)
    after.add_run().font.size = Pt(2)


def add_note(document: Document, text: str) -> None:
    table = document.add_table(rows=1, cols=1)
    cell = table.cell(0, 0)
    set_cell_shading(cell, LIGHT_BLUE)
    set_cell_margins(cell, top=120, start=170, bottom=120, end=170)
    prevent_row_split(table.rows[0])
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(0)
    add_inline(paragraph, text)
    for run in paragraph.runs:
        run.font.size = Pt(9.2)
    document.add_paragraph().paragraph_format.space_after = Pt(0)


def add_image(
    document: Document,
    source_path: Path,
    alt_text: str,
    width_value: float = 4.0,
    width_unit: str = "in",
) -> None:
    if source_path.suffix.lower() == ".svg":
        png_path = source_path.with_suffix(".png")
        if not png_path.is_file():
            raise FileNotFoundError(f"SVG needs a PNG sidecar for DOCX: {png_path}")
        source_path = png_path
    if not source_path.is_file():
        raise FileNotFoundError(f"Missing image: {source_path}")
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.keep_with_next = True
    section = document.sections[-1]
    usable_width = section.page_width - section.left_margin - section.right_margin
    desired_width = Inches(width_value) if width_unit == "in" else Cm(width_value)
    shape = paragraph.add_run().add_picture(str(source_path), width=min(desired_width, usable_width))
    max_height = int((section.page_height - section.top_margin - section.bottom_margin) * 0.80)
    if shape.height > max_height:
        shape.width = int(shape.width * max_height / shape.height)
        shape.height = max_height
    shape._inline.docPr.set("descr", alt_text)
    shape._inline.docPr.set("title", alt_text)
    caption = document.add_paragraph(style="Caption")
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.paragraph_format.space_after = Pt(6)
    caption.add_run(alt_text)


def markdown_to_docx(
    document: Document, source: Path, *, text: str | None = None,
    allow_missing_images: bool = False,
) -> None:
    lines = (text if text is not None else source.read_text(encoding="utf-8")).splitlines()
    first_heading = True
    index = 0
    ordered_count = 0
    in_ordered_list = False
    page_break_pending = False
    while index < len(lines):
        stripped = lines[index].strip()
        if not stripped:
            index += 1
            continue
        if stripped == "<!-- pagebreak -->":
            # This explicit author-controlled marker is preserved in Markdown,
            # but renders as a page break rather than literal HTML.
            page_break_pending = True
            in_ordered_list = False
            index += 1
            continue
        if stripped.startswith("<!--"):
            while index < len(lines) and "-->" not in lines[index]:
                index += 1
            index += 1
            continue
        if re.fullmatch(r"<br\s*/?>", stripped, flags=re.IGNORECASE):
            document.add_paragraph()
            index += 1
            continue
        if stripped.startswith("```"):
            fence_line = index + 1
            index += 1
            code_lines = []
            while index < len(lines) and not lines[index].strip().startswith("```"):
                code_lines.append(lines[index])
                index += 1
            if index >= len(lines):
                raise ValueError(f"Unclosed code block starting at line {fence_line}")
            add_code_block(document, "\n".join(code_lines))
            in_ordered_list = False
            index += 1
            continue
        image_match = re.fullmatch(
            r"!\[([^\]]*)\]\(([^)]+)\)(?:\s*\{width=([0-9]+(?:\.[0-9]+)?)(in|cm)\})?",
            stripped, flags=re.IGNORECASE,
        )
        if image_match:
            image_path = (source.parent / image_match.group(2)).resolve()
            try:
                add_image(document, image_path, image_match.group(1),
                          float(image_match.group(3) or 4.0),
                          (image_match.group(4) or "in").lower())
            except FileNotFoundError:
                if not allow_missing_images:
                    raise
                print(f"WARNING: image omitted: {image_path}", file=sys.stderr)
                add_note(document, f"[DRAFT: missing image {image_match.group(2)}]")
            in_ordered_list = False
            index += 1
            continue
        if stripped.startswith("|"):
            table_lines = []
            while index < len(lines) and lines[index].strip().startswith("|"):
                table_lines.append(lines[index])
                index += 1
            add_markdown_table(document, table_lines)
            in_ordered_list = False
            continue
        if stripped.startswith(">"):
            note_lines = []
            while index < len(lines) and lines[index].strip().startswith(">"):
                note_lines.append(lines[index].strip()[1:].strip())
                index += 1
            if note_lines and re.fullmatch(r"\[!NOTE\]", note_lines[0], flags=re.IGNORECASE):
                note_lines.pop(0)
            add_note(document, " ".join(note_lines))
            in_ordered_list = False
            continue
        heading_match = re.match(r"^(#{1,6})\s+(.+)$", stripped)
        if heading_match:
            level = len(heading_match.group(1))
            style = "Title" if first_heading else f"Heading {level}"
            paragraph = document.add_paragraph(style=style)
            if page_break_pending:
                paragraph.paragraph_format.page_break_before = True
                page_break_pending = False
            add_inline(paragraph, heading_match.group(2), base_bold=level == 2)
            first_heading = False
            in_ordered_list = False
            index += 1
            continue
        checkbox = re.match(r"^- \[([ xX])\] (.*)$", stripped)
        if checkbox:
            paragraph = document.add_paragraph(style="List Bullet")
            mark = "☒" if checkbox.group(1).lower() == "x" else "☐"
            add_inline(paragraph, f"{mark} {checkbox.group(2)}")
            in_ordered_list = False
            index += 1
            continue
        if re.match(r"^[-*]\s+", stripped):
            paragraph = document.add_paragraph(style="List Bullet")
            add_inline(paragraph, re.sub(r"^[-*]\s+", "", stripped))
            in_ordered_list = False
            index += 1
            continue
        ordered = re.match(r"^(\d+)\.\s+(.*)$", stripped)
        if ordered:
            number = int(ordered.group(1))
            ordered_count = (ordered_count + 1) if in_ordered_list and number == 1 else number
            paragraph = document.add_paragraph(style="Ordered List")
            add_inline(paragraph, f"{ordered_count}. {ordered.group(2)}")
            in_ordered_list = True
            index += 1
            continue
        paragraph_style = "Document Subtitle" if "STR Integration Documentation" in stripped and index < 5 else "Normal"
        paragraph = document.add_paragraph(style=paragraph_style)
        if page_break_pending:
            paragraph.paragraph_format.page_break_before = True
            page_break_pending = False
        add_inline(paragraph, stripped)
        in_ordered_list = False
        index += 1


def add_document_end_mark(document: Document) -> None:
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(10)
    run = paragraph.add_run("— End of guide —")
    run.italic = True
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor.from_string(MID_BLUE)


def sanitize_extended_properties(path: Path, title: str) -> None:
    """Replace stale template title/statistics in docProps/app.xml."""
    ep_ns = "http://schemas.openxmlformats.org/officeDocument/2006/extended-properties"
    vt_ns = "http://schemas.openxmlformats.org/officeDocument/2006/docPropsVTypes"
    ElementTree.register_namespace("", ep_ns)
    ElementTree.register_namespace("vt", vt_ns)

    with zipfile.ZipFile(path, "r") as source_zip:
        package = {item.filename: source_zip.read(item.filename) for item in source_zip.infolist()}

    app_path = "docProps/app.xml"
    if app_path not in package:
        return
    root = ElementTree.fromstring(package[app_path])
    for name in ("Company", "Manager", "Template"):
        node = root.find(f"{{{ep_ns}}}{name}")
        if node is not None:
            node.text = ""
    titles = root.find(f"{{{ep_ns}}}TitlesOfParts")
    if titles is not None:
        for value in titles.iter(f"{{{vt_ns}}}lpstr"):
            value.text = title
    for property_name in ("Pages", "Words", "Characters", "Lines", "Paragraphs", "CharactersWithSpaces"):
        node = root.find(f"{{{ep_ns}}}{property_name}")
        if node is not None:
            node.text = "0"
    package[app_path] = ElementTree.tostring(root, encoding="utf-8", xml_declaration=True)

    with tempfile.NamedTemporaryFile(
        dir=path.parent, prefix=f".{path.stem}.", suffix=".docx", delete=False
    ) as temporary:
        temporary_path = Path(temporary.name)
    try:
        with zipfile.ZipFile(temporary_path, "w", compression=zipfile.ZIP_DEFLATED) as target_zip:
            for name, data in package.items():
                target_zip.writestr(name, data)
        temporary_path.replace(path)
    finally:
        temporary_path.unlink(missing_ok=True)


def select_section(text: str, heading: str) -> str:
    """Select one exact heading and its children; ignore headings inside fences."""
    lines = text.splitlines()
    start: int | None = None
    level = 0
    in_code = False
    end = len(lines)
    for index, line in enumerate(lines):
        stripped = line.strip()
        if stripped.startswith("```"):
            in_code = not in_code
            continue
        if in_code:
            continue
        match = re.match(r"^(#{1,6})\s+(.+)$", stripped)
        if not match:
            continue
        if start is None and match.group(2).strip() == heading:
            start, level = index, len(match.group(1))
        elif start is not None and len(match.group(1)) <= level:
            end = index
            break
    if start is None:
        raise ValueError(f"Section not found: {heading}")
    selected = lines[start:end]
    while selected and selected[-1].strip() in ("", "<!-- pagebreak -->"):
        selected.pop()
    return "\n".join(selected) + "\n"


def infer_title(text: str) -> str:
    for line in text.splitlines():
        match = re.match(r"^#{1,6}\s+(.+)$", line.strip())
        if match:
            return match.group(1).strip()
    raise ValueError("Markdown must contain a document heading.")


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--source", type=Path, required=True, help="Authoritative Markdown source")
    parser.add_argument("--output", type=Path, help="Defaults to SOURCE.docx; specify for a section export")
    parser.add_argument("--template", type=Path, help="Optional .docx layout/style template")
    parser.add_argument("--title", help="Override the first Markdown heading and Word metadata")
    parser.add_argument("--footer", help="Default: inferred/overridden title followed by FIRE / STR")
    parser.add_argument("--subject", default="STR Integration Documentation")
    parser.add_argument("--author", default="FIRE / STR")
    parser.add_argument("--section", help="Export this exact heading and its children only")
    parser.add_argument("--write-markdown", type=Path, help="Also save selected/rendered Markdown")
    parser.add_argument("--allow-missing-images", action="store_true",
                        help="Draft only: show a warning instead of failing for a missing image")
    return parser.parse_args()


def main() -> int:
    args = parse_args()
    source = args.source.resolve()
    if not source.is_file():
        raise FileNotFoundError(source)
    if args.template is not None and not args.template.is_file():
        raise FileNotFoundError(args.template)
    if args.section and args.output is None:
        raise ValueError("Use --output for a section export to avoid replacing the full guide.")
    output = (args.output or source.with_suffix(".docx")).resolve()
    if output == source or output.suffix.lower() != ".docx":
        raise ValueError("Output must be a .docx path different from the source.")
    text = source.read_text(encoding="utf-8")
    if args.section:
        text = select_section(text, args.section)
    title = args.title or infer_title(text)
    if args.title:
        text = re.sub(r"^(#{1,6})\s+.+$", lambda m: m.group(1) + " " + title,
                      text, count=1, flags=re.MULTILINE)
    metadata = GuideMetadata(title, args.footer or f"{title} | FIRE / STR",
                             args.subject, args.author)
    document = Document(args.template) if args.template else Document()
    if args.template is None:
        section = document.sections[0]
        section.page_width, section.page_height = Inches(8.5), Inches(11)
        section.left_margin = section.right_margin = Inches(0.6)
        section.top_margin = section.bottom_margin = Inches(0.6)
        section.header_distance = section.footer_distance = Inches(0.25)
    sanitize_template(document, metadata)
    configure_styles(document)
    markdown_to_docx(document, source, text=text, allow_missing_images=args.allow_missing_images)
    add_document_end_mark(document)
    output.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.NamedTemporaryFile(dir=output.parent, suffix=".docx", delete=False) as handle:
        temporary = Path(handle.name)
    try:
        document.save(temporary)
        sanitize_extended_properties(temporary, title)
        temporary.replace(output)
    finally:
        temporary.unlink(missing_ok=True)
    if args.write_markdown:
        md_output = args.write_markdown.resolve()
        if md_output == source:
            raise ValueError("--write-markdown must not overwrite the authoritative source.")
        md_output.parent.mkdir(parents=True, exist_ok=True)
        md_output.write_text(text, encoding="utf-8")
    print(output)
    return 0


if __name__ == "__main__":
    try:
        raise SystemExit(main())
    except (FileNotFoundError, ValueError) as exc:
        print(f"ERROR: {exc}", file=sys.stderr)
        raise SystemExit(2)


